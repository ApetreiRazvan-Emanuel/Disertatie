# -*- coding: utf-8 -*-
"""Approximate PIL renderer for the variant pptx files — layout QA only.
Usage: python render_variant.py <deck.pptx> <outdir>"""
import os, io, sys
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.util import Emu
from pptx.enum.text import PP_ALIGN

PPTX = sys.argv[1]
OUT = sys.argv[2]
os.makedirs(OUT, exist_ok=True)
SCALE = 150.0

FONTS = {}
def font(size_pt, bold=False, italic=False):
    key = (round(size_pt), bold, italic)
    if key not in FONTS:
        name = "arialbd.ttf" if bold else ("ariali.ttf" if italic else "arial.ttf")
        if bold and italic: name = "arialbi.ttf"
        px = max(8, int(size_pt * SCALE / 72.0))
        FONTS[key] = ImageFont.truetype(os.path.join(r"C:\Windows\Fonts", name), px)
    return FONTS[key]

def emu2px(v):
    return int(Emu(v).inches * SCALE)

def color_of(color_format, default=(198, 207, 223)):
    try:
        if color_format and color_format.type is not None:
            rgb = color_format.rgb
            return (rgb[0], rgb[1], rgb[2])
    except Exception:
        pass
    return default

def wrap(draw, text, fnt, maxw):
    words = text.split()
    lines, cur = [], ""
    for w in words:
        t = (cur + " " + w).strip()
        if draw.textlength(t, font=fnt) <= maxw or not cur:
            cur = t
        else:
            lines.append(cur); cur = w
    if cur: lines.append(cur)
    return lines or [""]

def has_bullet(p):
    from pptx.oxml.ns import qn
    pPr = p._p.find(qn('a:pPr'))
    if pPr is None: return False
    return pPr.find(qn('a:buChar')) is not None

def draw_textframe(draw, tf, x, y, w, h):
    cy = y + 2
    for p in tf.paragraphs:
        if not p.runs:
            cy += 6
            continue
        r0 = p.runs[0]
        size = r0.font.size.pt if r0.font.size else 18
        fnt = font(size, bool(r0.font.bold), bool(r0.font.italic))
        col = color_of(r0.font.color)
        text = "".join(r.text for r in p.runs)
        bullet = has_bullet(p)
        indent = int(0.19 * SCALE) if bullet else 0
        lines = wrap(draw, text, fnt, w - indent - 4)
        lh = int(size * SCALE / 72.0 * 1.18)
        for li, line in enumerate(lines):
            tx = x + indent
            if p.alignment == PP_ALIGN.CENTER:
                tw = draw.textlength(line, font=fnt)
                tx = x + (w - tw) / 2
            elif p.alignment == PP_ALIGN.RIGHT:
                tw = draw.textlength(line, font=fnt)
                tx = x + w - tw
            if bullet and li == 0:
                draw.text((x + 2, cy), "•", font=fnt, fill=col)
            draw.text((tx, cy), line, font=fnt, fill=col)
            cy += lh
        sa = p.space_after.pt if p.space_after else 0
        cy += int(sa * SCALE / 72.0)
    return cy - y

prs = Presentation(PPTX)
SW, SH = emu2px(prs.slide_width), emu2px(prs.slide_height)

for idx, slide in enumerate(prs.slides, start=1):
    bg = (255, 255, 255)
    try:
        if slide.background.fill.type is not None:
            rgb = slide.background.fill.fore_color.rgb
            bg = (rgb[0], rgb[1], rgb[2])
    except Exception:
        pass
    img = Image.new("RGB", (SW, SH), bg)
    draw = ImageDraw.Draw(img)
    for shape in slide.shapes:
        x, y = emu2px(shape.left or 0), emu2px(shape.top or 0)
        try:
            w, h = emu2px(shape.width), emu2px(shape.height)
        except Exception:
            w = h = 10
        if shape.shape_type == 13 or shape.__class__.__name__ == "Picture":
            try:
                im = Image.open(io.BytesIO(shape.image.blob)).convert("RGBA")
                im = im.resize((max(1, w), max(1, h)))
                img.paste(im, (x, y), im)
            except Exception:
                draw.rectangle([x, y, x + w, y + h], outline="red")
        elif shape.has_table:
            from pptx.oxml.ns import qn as _qn
            tbl = shape.table
            col_w = [emu2px(c.width) for c in tbl.columns]
            row_h = [emu2px(r.height) for r in tbl.rows]
            xs = [x]
            for wpx in col_w:
                xs.append(xs[-1] + wpx)
            ysr = [y]
            for hpx in row_h:
                ysr.append(ysr[-1] + hpx)
            border_lines = []
            for ri in range(len(row_h)):
                for ci in range(len(col_w)):
                    cell = tbl.cell(ri, ci)
                    try:
                        if cell.is_spanned:
                            continue
                        sw, sh = cell.span_width or 1, cell.span_height or 1
                    except Exception:
                        sw = sh = 1
                    x0, y0 = xs[ci], ysr[ri]
                    x1 = xs[min(ci + sw, len(col_w))]
                    y1 = ysr[min(ri + sh, len(row_h))]
                    fill = bg
                    try:
                        if cell.fill.type is not None and cell.fill.type != 5:
                            rgb = cell.fill.fore_color.rgb
                            fill = (rgb[0], rgb[1], rgb[2])
                    except Exception:
                        pass
                    draw.rectangle([x0, y0, x1, y1], fill=fill)
                    tfp = cell.text_frame.paragraphs[0]
                    if tfp.runs:
                        r0 = tfp.runs[0]
                        size = r0.font.size.pt if r0.font.size else 12
                        fnt = font(size, bool(r0.font.bold))
                        col = color_of(r0.font.color)
                        text = "".join(r.text for r in tfp.runs)
                        tw = draw.textlength(text, font=fnt)
                        th = int(size * SCALE / 72.0 * 1.1)
                        tx = x0 + 6 if tfp.alignment == PP_ALIGN.LEFT else x0 + (x1 - x0 - tw) / 2
                        draw.text((tx, y0 + (y1 - y0 - th) / 2), text, font=fnt, fill=col)
                    tcPr = cell._tc.find(_qn("a:tcPr"))
                    if tcPr is None:
                        continue
                    for tag, seg in (("a:lnL", (x0, y0, x0, y1)), ("a:lnR", (x1, y0, x1, y1)),
                                     ("a:lnT", (x0, y0, x1, y0)), ("a:lnB", (x0, y1, x1, y1))):
                        ln = tcPr.find(_qn(tag))
                        if ln is None:
                            continue
                        wpt = int(ln.get("w", "12700")) / 12700.0
                        wpx = max(1, round(wpt * SCALE / 72.0))
                        clr = ln.find(_qn("a:solidFill") + "/" + _qn("a:srgbClr"))
                        lc = (255, 255, 255)
                        if clr is not None:
                            v = clr.get("val")
                            lc = tuple(int(v[k:k+2], 16) for k in (0, 2, 4))
                        border_lines.append((seg, wpx, lc))
            for seg, wpx, lc in border_lines:
                draw.line(seg, fill=lc, width=wpx)
        elif shape.has_text_frame and shape.shape_type == 17:
            used = draw_textframe(draw, shape.text_frame, x, y, w, h)
            if used > h + 8:
                draw.rectangle([x, y, x + w, y + used], outline="orange", width=2)
        else:
            fill = None
            try:
                if shape.fill.type is not None and shape.fill.type != 5:
                    rgb = shape.fill.fore_color.rgb
                    fill = (rgb[0], rgb[1], rgb[2])
            except Exception:
                pass
            line_col = None
            try:
                rgb = shape.line.color.rgb
                line_col = (rgb[0], rgb[1], rgb[2])
            except Exception:
                pass
            if fill or line_col:
                draw.rounded_rectangle([x, y, x + w, y + h], radius=8, fill=fill, outline=line_col)
            if shape.has_text_frame and shape.text_frame.text.strip():
                draw_textframe(draw, shape.text_frame, x + 6, y + 4, w - 12, h)
    img.save(os.path.join(OUT, f"slide{idx:02d}.png"))
print("rendered to", OUT)
